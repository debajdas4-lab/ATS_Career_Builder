from io import BytesIO
import json
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "13293F"
TEAL = "1F6F8B"
GOLD = "C2A14D"
BODY = "263238"
MUTED = "667781"

SECTION_HEADINGS = {
    "SUMMARY", "PROFESSIONAL SUMMARY", "CORE SKILLS", "SKILLS",
    "PROFESSIONAL EXPERIENCE", "EXPERIENCE", "EDUCATION", "CERTIFICATIONS", "PROJECTS",
}

PROFILE = (
    "Transformation Leader and Senior Program Manager with 24+ years of experience driving "
    "enterprise-wide digital transformation, business process optimization, operational excellence, "
    "and large-scale technology modernization initiatives. Proven track record leading global cross-functional "
    "teams, improving customer and employee experiences, establishing KPI-driven governance frameworks, "
    "managing strategic programs, and delivering measurable business outcomes. Skilled in process re-engineering, "
    "change management, risk mitigation, stakeholder engagement, financial governance, Agile delivery, and "
    "continuous improvement. Recognized for leading complex transformation programs across SAP S/4HANA, "
    "regulatory compliance, enterprise platforms, and operational modernization initiatives."
)

EXECUTIVE_EXPERTISE = [
    "Enterprise Transformation & Operations",
    "Process Optimization & Continuous Improvement",
    "Program Governance & KPI Reporting",
    "Digital Transformation & Change Management",
    "Risk Management & Mitigation",
    "Customer Experience Improvement",
    "Executive Stakeholder Leadership",
    "Agile Delivery & Cross-Functional Leadership",
]

CORE_COMPETENCIES = (
    "Enterprise Program Management | AI Platform Governance | Digital Transformation | SAP S/4HANA | "
    "Application Modernization | Agile | Azure DevOps | CI/CD | Governance, Risk & Financial Management | "
    "Stakeholder Leadership | Microsoft Copilot | Generative AI | Intelligent Automation"
)


def _set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_margins(cell, top=120, start=160, bottom=120, end=160):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def _run_font(run, name="Aptos", size=10.5, color=BODY, bold=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def _heading(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    _run_font(r, size=11, color=NAVY, bold=True)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "14")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), GOLD)
    borders.append(bottom)
    p_pr.append(borders)


def _body(doc, text: str, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    _run_font(p.add_run(text))


def _body_with_lead(doc, lead: str, text: str, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    _run_font(p.add_run(lead), bold=True)
    _run_font(p.add_run(text))


def _bullet(doc, text: str, compact=False):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(2 if compact else 4)
    p.paragraph_format.line_spacing = 1.03
    _run_font(p.add_run(text))


def _header(doc):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(7.1)
    cell = table.cell(0, 0)
    cell.width = Inches(7.1)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_shading(cell, NAVY)
    _set_cell_margins(cell, top=170, start=55, bottom=150, end=55)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("DEBA JYOTI DAS")
    _run_font(r, name="Aptos Display", size=25, color="FFFFFF", bold=True)

    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Transformation Leader | Enterprise Digital Transformation | Process Excellence | Business Operations | Operational\nEfficiency | Program Governance | SAP S/4HANA | Change Management | PMP")
    _run_font(r, size=10.5, color="EAD9A6", bold=True)

    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("debajdas4@gmail.com       ·       +91 96633 88311       ·       Bangalore, India       ·       linkedin.com/in/debajyoti-das")
    _run_font(r, size=9, color="DCE3EC")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def _extract_json_payload(text: str) -> dict | None:
    candidate = text.strip()
    start = candidate.find("{")
    if start < 0:
        return None
    candidate = candidate[start:]
    try:
        payload, _ = json.JSONDecoder().raw_decode(candidate)
        return payload if isinstance(payload, dict) else None
    except json.JSONDecodeError:
        match = re.search(r'"optimized_resume"\s*:\s*"(.*?)(?:"\s*,\s*"cover_note"|"\s*})', candidate, flags=re.DOTALL)
        if match:
            value = match.group(1).replace("\\n", "\n").replace('\\"', '"')
            return {"optimized_resume": value}
        return None


def clean_resume_text(text: str) -> str:
    cleaned = text.strip()
    cleaned = re.sub(r"^\s*```(?:json|text|markdown)?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s*```\s*$", "", cleaned)
    payload = _extract_json_payload(cleaned)
    if payload and payload.get("optimized_resume"):
        cleaned = str(payload["optimized_resume"])
    return cleaned.strip()


def _content_lines(text: str) -> list[str]:
    lines = [line.strip() for line in clean_resume_text(text).splitlines() if line.strip()]
    result = []
    active = False
    for line in lines:
        normalized = line.upper().rstrip(":")
        if normalized in {"PROFESSIONAL EXPERIENCE", "EXPERIENCE"}:
            active = True
            continue
        if normalized in {"EDUCATION", "CERTIFICATIONS", "EDUCATION AND CERTIFICATION", "CORE COMPETENCIES"}:
            active = False
            continue
        if normalized in {"SUMMARY", "PROFESSIONAL SUMMARY", "CORE SKILLS", "SKILLS", "ATS-OPTIMIZED RESUME DRAFT", "RESUME"}:
            continue
        if active:
            result.append(line)
    return result


def build_docx(resume_text: str) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    doc.core_properties.title = "Deba Jyoti Das - ATS Optimized Resume"

    _header(doc)
    _heading(doc, "Profile")
    lead = "Transformation Leader and Senior Program Manager"
    _body_with_lead(doc, lead, PROFILE[len(lead):])
    _heading(doc, "Executive Expertise")
    for item in EXECUTIVE_EXPERTISE:
        _bullet(doc, item, compact=True)

    _heading(doc, "Professional Experience")
    dynamic = _content_lines(resume_text)
    for line in dynamic:
        if line.startswith(("- ", "* ", "• ")):
            _bullet(doc, line[2:].strip())
        elif re.match(r"^(Recognition\.|[A-Z][A-Za-z .&'-]+ (?:at|\|))", line):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            _run_font(p.add_run(line), size=10.5, color=NAVY, bold=True)
        else:
            _bullet(doc, line)

    _heading(doc, "Core Competencies")
    _body(doc, CORE_COMPETENCIES)
    _heading(doc, "Education and Certification")
    _body(doc, "Master of Engineering (Honors) | National Institute of Technology (NIT), Tiruchirappalli | CGPA: 8.6", after=3)
    _body(doc, "Bachelor of Engineering (Honors) | Assam Engineering College | 75 percent", after=3)
    _body(doc, "GATE: 88.0 percentile | PMP: Project Management Professional | Gen AI / Agentic AI Certification", after=3)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run_font(footer.add_run("ATS-friendly draft | Verify all details before submission"), size=8, color=MUTED)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
